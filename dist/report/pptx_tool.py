#!/usr/bin/env python
"""Peekly PPT report generator — python-pptx subprocess worker.

Invoked by the Node MCP tool `generate_report` (see ../tools/report.ts) via
child_process. Reads one JSON job from stdin and writes one JSON result
line to stdout; never writes anything else to stdout (diagnostics, if any,
must go to stderr) so the caller can parse stdout as a single JSON value.

stdin:
  {
    "templatePath": "<absolute path to read-only .pptx template>",
    "outputPath": "<absolute path to write the generated .pptx to>",
    "issues": [
      {
        "breadcrumb": str, "screenshotPath": str, "problem": str, "improvement": str,
        "problemArea": {  # optional
          "xFraction": float, "yFraction": float,   # top-left, 0-1 of screenshot width/height
          "widthFraction": float, "heightFraction": float
        }
      },
      ...
    ]
  }

stdout (exactly one line, success):
  {"success": true, "outputPath": "...", "slideCount": N}

stdout (exactly one line, failure):
  {"success": false, "error": "..."}

------------------------------------------------------------------------
PROVISIONAL / UNCONFIRMED — see CONVENTIONS.md section 8 and design doc
section 3.6. No real PPT template exists yet (the user will supply one
later). Until then this script assumes a placeholder-marker convention
invented for this MVP, which WILL need re-validation against the real
template:

  - Slide 0 of the template is the cover slide (untouched).
  - Slide index 1 is the "issue slide" template. It is cloned once per
    confirmed issue, then the original template slide is dropped, so the
    final deck is: cover + one slide per issue.
  - Text fields are substituted via literal token replacement inside text
    frames: {{breadcrumb}}, {{problem}}, {{improvement}}. Replacement only
    works if a whole token lives inside a single text run — if a real
    template's authoring tool splits "{{breadcrumb}}" across multiple runs
    (e.g. due to autocorrect/spell-check boundaries), the token will not
    be found. This is a known limitation of the run-level find/replace
    technique, not something this script works around.
  - The screenshot field is substituted by locating a Picture shape whose
    `shape.name == "screenshot"` on the issue-slide template and swapping
    its embedded image. If no such shape exists, this is treated as a hard
    error for that issue (never silently skipped) — a QA report silently
    missing a screenshot would be a serious, easy-to-miss defect.

None of the above should be treated as a confirmed spec; it is a best
guess to unblock the MVP pipeline pending the real template file.
------------------------------------------------------------------------

Two fixes from real-usage feedback (a real generated report opened in
PowerPoint showed the screenshot partially covered by overflowing text):

  - Text placeholders no longer grow to fit long real problem/improvement
    text (which risked overlapping the screenshot below/near them) — see
    `_replace_text_tokens`, which instead sets
    MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE so PowerPoint shrinks the font to fit
    the box's ORIGINAL size, same as its own "Shrink text on overflow"
    option.
  - Each issue may optionally include a `problemArea` (fractions 0-1 of
    the screenshot's own width/height, not pixels) marking exactly where
    the problem is; when present, `_draw_problem_area` draws a no-fill,
    outline-only rectangle over that part of the screenshot so it never
    obscures the image underneath it.
------------------------------------------------------------------------

Slide-duplication technique: python-pptx has no native "duplicate slide"
API. This uses the well-known workaround of deep-copying the template
slide's shape XML onto a freshly added blank-layout slide, while also
copying (and rId-remapping) the template slide's non-layout relationships
— chiefly the image relationship backing its placeholder picture — so
that the copied shapes' r:embed/r:id references keep resolving correctly
in the new slide part.

`_remove_slide_at` drops both the sldIdLst entry and the presentation
part's relationship to the slide part, so the removed template slide's
XML part becomes unreachable and is excluded by Package.iter_parts() on
save — see its docstring. An earlier version only removed the sldIdLst
entry, which left the part reachable and caused the next add_slide() to
reuse its partname (e.g. slide3.xml), producing two same-named zip
entries and a PowerPoint "repair" prompt on open. Confirmed via a real
generated report opened in PowerPoint.
"""
from __future__ import annotations

import copy
import json
import os
import shutil
import sys
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from pptx.util import Pt

# PROVISIONAL constants — see module docstring.
ISSUE_TEMPLATE_SLIDE_INDEX = 1
SCREENSHOT_SHAPE_NAME = "screenshot"
PROBLEM_AREA_LINE_COLOR = RGBColor(0xFF, 0x00, 0x00)
PROBLEM_AREA_LINE_WIDTH = Pt(2.25)
# Extra breathing room around the tight element bounds so the outline
# doesn't hug the flagged element pixel-for-pixel (purely cosmetic margin).
PROBLEM_AREA_PADDING = Pt(6)

_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_R_NS_PREFIX = "{%s}" % _R_NS


def _duplicate_slide(prs: Presentation, template_slide) -> Any:
    """Append a new slide that is a deep copy of `template_slide` (shapes +
    non-layout relationships) and return it."""
    layout = template_slide.slide_layout
    new_slide = prs.slides.add_slide(layout)

    # add_slide() auto-populates placeholder shapes from the layout; discard
    # them — we want an exact copy of the template slide's own shapes.
    for shape in list(new_slide.shapes):
        shape._element.getparent().remove(shape._element)

    # Copy relationships (images, hyperlinks, ...), excluding the slide's
    # link to its layout (add_slide() already established that one).
    # get_or_add()/get_or_add_ext_rel() mint a fresh rId in the new slide
    # part, so track an old->new map to rewrite references inside the
    # copied shape XML below.
    rid_remap: dict[str, str] = {}
    for rid, rel in template_slide.part.rels.items():
        if rel.reltype == RT.SLIDE_LAYOUT:
            continue
        if rel.is_external:
            new_rid = new_slide.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new_rid = new_slide.part.rels.get_or_add(rel.reltype, rel.target_part)
        if new_rid != rid:
            rid_remap[rid] = new_rid

    for shape in template_slide.shapes:
        new_el = copy.deepcopy(shape._element)
        if rid_remap:
            for el in new_el.iter():
                for attr_name, attr_val in list(el.attrib.items()):
                    if attr_name.startswith(_R_NS_PREFIX) and attr_val in rid_remap:
                        el.set(attr_name, rid_remap[attr_val])
        new_slide.shapes._spTree.insert_element_before(new_el, "p:extLst")

    return new_slide


def _remove_slide_at(prs: Presentation, index: int) -> None:
    """Fully drop the slide currently at `index`: remove it from the slide
    order AND drop the presentation part's relationship to its slide part.

    Removing only the sldIdLst entry (the once-documented "acceptable"
    shortcut) leaves the slide part reachable via prs.part.rels, so
    Package.iter_parts()/next_partname() still counts it. The next
    add_slide() then picks the same partname (e.g. slide3.xml) for a new
    slide, and both the orphaned original and the new slide get serialized
    under that identical zip entry name on save — which PowerPoint detects
    as corruption and offers to "repair" on open. Dropping the relationship
    here makes the old part unreachable so it is correctly excluded from
    iter_parts() and never written.
    """
    id_list = prs.slides._sldIdLst
    sld_ids = list(id_list)
    sld_id = sld_ids[index]
    rId = sld_id.rId
    id_list.remove(sld_id)
    prs.part.rels.pop(rId)


def _replace_text_tokens(slide, mapping: dict[str, str]) -> None:
    """Substitute placeholder tokens in-place, then force each affected
    text box to shrink its font to fit its EXISTING size rather than
    growing the box — a long real problem/improvement description must
    not resize its placeholder into overlapping neighboring shapes (e.g.
    the screenshot). This mirrors PowerPoint's own "Shrink text on
    overflow" autofit option, applied via python-pptx's auto_size enum."""
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        replaced = False
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                for token, value in mapping.items():
                    if token in run.text:
                        run.text = run.text.replace(token, value)
                        replaced = True
        if replaced:
            shape.text_frame.word_wrap = True
            shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _replace_screenshot(slide, image_path: str):
    """Swap the embedded image on the `screenshot`-named Picture shape and
    return that shape, so the caller can position a problem-area
    annotation relative to it."""
    shape = None
    for candidate in slide.shapes:
        if candidate.name == SCREENSHOT_SHAPE_NAME:
            shape = candidate
            break
    if shape is None:
        raise LookupError(
            f"no Picture shape named '{SCREENSHOT_SHAPE_NAME}' found on the issue-slide template "
            "(PROVISIONAL placeholder-marker convention — see this script's module docstring; "
            "re-validate against the real template)"
        )
    if not os.path.isfile(image_path):
        raise FileNotFoundError(image_path)

    blip = shape._element.blipFill.blip
    if blip is None:
        raise LookupError(f"shape '{SCREENSHOT_SHAPE_NAME}' has no blip fill (not a picture shape?)")

    _, rid = slide.part.get_or_add_image_part(image_path)
    blip.set(qn("r:embed"), rid)
    return shape


def _draw_problem_area(slide, screenshot_shape, problem_area: dict[str, float]) -> None:
    """Draw a no-fill, outlined-only rectangle over the given fraction
    (0-1) of the screenshot shape's own bounds, to point out exactly
    where on the screenshot the problem is."""
    left = screenshot_shape.left + int(problem_area["xFraction"] * screenshot_shape.width)
    top = screenshot_shape.top + int(problem_area["yFraction"] * screenshot_shape.height)
    width = int(problem_area["widthFraction"] * screenshot_shape.width)
    height = int(problem_area["heightFraction"] * screenshot_shape.height)

    # Pad outward, then clamp to the screenshot's own bounds so the
    # outline never spills outside the image it's annotating.
    screenshot_right = screenshot_shape.left + screenshot_shape.width
    screenshot_bottom = screenshot_shape.top + screenshot_shape.height
    padded_left = max(screenshot_shape.left, left - PROBLEM_AREA_PADDING)
    padded_top = max(screenshot_shape.top, top - PROBLEM_AREA_PADDING)
    padded_right = min(screenshot_right, left + width + PROBLEM_AREA_PADDING)
    padded_bottom = min(screenshot_bottom, top + height + PROBLEM_AREA_PADDING)
    left, top = padded_left, padded_top
    width, height = padded_right - padded_left, padded_bottom - padded_top

    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.background()  # no fill — outline only, so it doesn't cover the screenshot underneath
    box.line.color.rgb = PROBLEM_AREA_LINE_COLOR
    box.line.width = PROBLEM_AREA_LINE_WIDTH
    box.shadow.inherit = False


def _generate_report(job: dict[str, Any]) -> dict[str, Any]:
    template_path = job["templatePath"]
    output_path = job["outputPath"]
    issues = job.get("issues", [])

    if not os.path.isfile(template_path):
        return {"success": False, "error": f"template not found: {template_path}"}

    # Original template is always read-only: copy first, then only ever
    # open/modify/save the copy at output_path. Never open template_path
    # itself with Presentation() for editing.
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    shutil.copyfile(template_path, output_path)

    prs = Presentation(output_path)

    if len(prs.slides) <= ISSUE_TEMPLATE_SLIDE_INDEX:
        return {
            "success": False,
            "error": (
                f"template must have at least {ISSUE_TEMPLATE_SLIDE_INDEX + 1} slides "
                f"(cover + issue template expected at index {ISSUE_TEMPLATE_SLIDE_INDEX}), "
                f"found {len(prs.slides)}"
            ),
        }

    template_slide = prs.slides[ISSUE_TEMPLATE_SLIDE_INDEX]

    for i, issue in enumerate(issues):
        for field in ("breadcrumb", "screenshotPath", "problem", "improvement"):
            if field not in issue:
                return {"success": False, "error": f"issues[{i}] missing required field '{field}'"}

        new_slide = _duplicate_slide(prs, template_slide)
        _replace_text_tokens(
            new_slide,
            {
                "{{breadcrumb}}": issue["breadcrumb"],
                "{{problem}}": issue["problem"],
                "{{improvement}}": issue["improvement"],
            },
        )
        try:
            screenshot_shape = _replace_screenshot(new_slide, issue["screenshotPath"])
        except FileNotFoundError:
            return {"success": False, "error": f"issues[{i}]: screenshot file not found: {issue['screenshotPath']}"}
        except LookupError as e:
            return {"success": False, "error": f"issues[{i}]: {e}"}

        problem_area = issue.get("problemArea")
        if problem_area is not None:
            _draw_problem_area(new_slide, screenshot_shape, problem_area)

    _remove_slide_at(prs, ISSUE_TEMPLATE_SLIDE_INDEX)

    slide_count = len(prs.slides)
    prs.save(output_path)

    return {"success": True, "outputPath": output_path, "slideCount": slide_count}


def main() -> int:
    try:
        # Read stdin as raw bytes and decode as UTF-8 explicitly — sys.stdin's
        # default text-mode encoding follows the OS locale (e.g. cp949 on
        # Korean Windows), which silently mangles the UTF-8 bytes Node writes
        # for non-ASCII (Korean) issue text into invalid lone surrogates.
        raw = sys.stdin.buffer.read().decode("utf-8")
        job = json.loads(raw)
    except json.JSONDecodeError as e:
        print(json.dumps({"success": False, "error": f"invalid JSON on stdin: {e}"}))
        return 1

    try:
        result = _generate_report(job)
    except Exception as e:  # top-level worker boundary: always report structured JSON, never a bare traceback on stdout
        print(json.dumps({"success": False, "error": f"{type(e).__name__}: {e}"}))
        return 1

    print(json.dumps(result))
    return 0 if result.get("success") else 1


if __name__ == "__main__":
    sys.exit(main())
